#!/usr/bin/env python3
"""Stefan Sense Bundle - local document helper (DOCX, PDF, XLSX, PPTX).

  python scripts/docs.py doctor
  python scripts/docs.py create spec.json out.docx|out.pdf|out.xlsx|out.pptx
  python scripts/docs.py inspect file.docx|file.pdf|file.xlsx|file.pptx
  python scripts/docs.py merge a.pdf b.pdf [...] out.pdf
  python scripts/docs.py split in.pdf out_dir

Independent implementation by M. Stefan Kassem (Stefan Sense), MIT licence. No network access.
Optional libraries: python-docx, reportlab, pypdf, openpyxl, python-pptx (probe with `doctor`).
"""
import argparse, html, importlib.util, json, os, sys, math, re, tempfile
from runtime_probe import probe_module
from pathlib import Path

LIBS = {'docx': 'docx', 'pdf': 'reportlab', 'pdf-read': 'pypdf', 'xlsx': 'openpyxl', 'pptx': 'pptx'}
FONT_CANDIDATES = [
    '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', '/usr/share/fonts/dejavu/DejaVuSans.ttf',
    '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf', '/usr/share/fonts/liberation-sans/LiberationSans-Regular.ttf',
    '/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf', '/usr/share/fonts/noto/NotoSans-Regular.ttf',
    '/Library/Fonts/Arial Unicode.ttf', '/System/Library/Fonts/Supplemental/Arial.ttf',
    'C:/Windows/Fonts/arial.ttf', 'C:/Windows/Fonts/calibri.ttf']


class SpecError(ValueError):
    pass


def need(module, pip_name):
    try:
        __import__(module)
    except (ImportError, OSError) as e:
        raise SpecError(f'{pip_name} cannot be loaded: {e}; use the host file tools or install the dependency in an authorized environment') from e


def object_spec(spec):
    if not isinstance(spec, dict): raise SpecError('spec must be a JSON object')
    return spec


def array(value, label):
    if not isinstance(value, list): raise SpecError(label + ' must be a JSON array')
    return value


def positive(value, label):
    if isinstance(value, bool): raise SpecError(label + ' must be a positive number')
    try: result = float(value)
    except (TypeError, ValueError): raise SpecError(label + ' must be a positive number')
    if not math.isfinite(result) or result <= 0: raise SpecError(label + ' must be finite and positive')
    return result


def table_spec(table):
    headers = array(table.get('headers'), 'headers')
    rows = array(table.get('rows', []), 'rows')
    if not headers or any(not isinstance(r, list) or len(r) != len(headers) for r in rows):
        raise SpecError('table rows must have the same number of cells as nonempty headers')
    table['rows'] = rows


def blocks_of(spec):
    """Normalize the document spec into a list of blocks (also accepts the legacy paragraphs/sections/tables form)."""
    if not isinstance(spec, dict):
        raise SpecError('spec must be a JSON object')
    blocks = [dict(b) if isinstance(b, dict) else b for b in array(spec.get('blocks', []), 'blocks')]
    for p in array(spec.get('paragraphs', []), 'paragraphs'):
        blocks.append({'type': 'paragraph', 'text': p})
    for s in array(spec.get('sections', []), 'sections'):
        if not isinstance(s, dict): raise SpecError('each section must be an object')
        array(s.get('paragraphs', []), 'section paragraphs')
        blocks.append({'type': 'heading', 'text': s.get('heading', ''), 'level': 1})
        blocks += [{'type': 'paragraph', 'text': p} for p in s.get('paragraphs', [])]
    for t in array(spec.get('tables', []), 'tables'):
        if not isinstance(t, dict): raise SpecError('each table must be an object')
        blocks.append(dict(t, type='table'))
    for b in blocks:
        if not isinstance(b, dict): raise SpecError('each block must be an object')
        kind = b.get('type')
        if kind in ('heading', 'paragraph') and 'text' not in b: raise SpecError(kind + ' needs text')
        if kind in ('bullets', 'numbered'): array(b.get('items', []), 'items')
        if kind == 'image':
            if not isinstance(b.get('path'), str) or not Path(b['path']).is_file(): raise SpecError('image path must name an existing file')
            for key in ('width_in', 'height_in'):
                if key in b: positive(b[key], key)
        if kind not in ('heading', 'paragraph', 'bullets', 'numbered', 'table', 'page_break', 'image'):
            raise SpecError(f'unknown block type: {kind}')
        if kind == 'table':
            table_spec(b)
    return blocks


# ---------------------------------------------------------------- DOCX
def create_docx(spec, out):
    object_spec(spec)
    need('docx', 'python-docx')
    from docx import Document
    from docx.shared import Pt, Inches
    d = Document()
    from docx.shared import Mm
    for section in d.sections:
        if str(spec.get('page_size', 'A4')).upper() == 'LETTER':
            section.page_width, section.page_height = Inches(8.5), Inches(11)
        else: section.page_width, section.page_height = Mm(210), Mm(297)
    style = d.styles['Normal']; style.font.name = spec.get('font', 'Calibri'); style.font.size = Pt(positive(spec.get('font_size', 11), 'font_size'))
    if spec.get('title'): d.add_heading(str(spec['title']), 0)
    if spec.get('subtitle'): d.add_paragraph().add_run(str(spec['subtitle'])).italic = True
    if spec.get('author'): d.core_properties.author = str(spec['author'])
    if spec.get('title'): d.core_properties.title = str(spec['title'])
    for b in blocks_of(spec):
        t = b['type']
        if t == 'heading': d.add_heading(str(b['text']), max(1, min(int(b.get('level', 1)), 4)))
        elif t == 'paragraph': d.add_paragraph(str(b['text']))
        elif t in ('bullets', 'numbered'):
            for item in b.get('items', []):
                d.add_paragraph(str(item), style='List Bullet' if t == 'bullets' else 'List Number')
        elif t == 'table':
            tab = d.add_table(rows=1, cols=len(b['headers'])); tab.style = 'Table Grid'
            for c, v in zip(tab.rows[0].cells, b['headers']):
                c.text = str(v)
                for r in c.paragraphs[0].runs: r.bold = True
            for row in b['rows']:
                for c, v in zip(tab.add_row().cells, row): c.text = '' if v is None else str(v)
        elif t == 'page_break': d.add_page_break()
        elif t == 'image':
            if not Path(b['path']).is_file(): raise SpecError('image not found: ' + b['path'])
            d.add_picture(b['path'], width=Inches(float(b.get('width_in', 6))))
    d.save(out)


def inspect_docx(src):
    need('docx', 'python-docx')
    from docx import Document
    d = Document(src)
    paras = [{'style': p.style.name, 'text': p.text} for p in d.paragraphs if p.text.strip()]
    return {'format': 'docx', 'headings': [p['text'] for p in paras if p['style'].lower().startswith(('heading', 'title'))],
            'paragraphs': paras, 'tables': [[[c.text for c in r.cells] for r in t.rows] for t in d.tables],
            'core_properties': {'title': d.core_properties.title, 'author': d.core_properties.author}}


# ---------------------------------------------------------------- PDF
def find_font(spec):
    path = spec.get('font_path')
    if path:
        if not Path(path).is_file(): raise SpecError('font_path not found: ' + path)
        return path
    return next((p for p in FONT_CANDIDATES if Path(p).is_file()), None)


def create_pdf(spec, out):
    object_spec(spec)
    need('reportlab', 'reportlab')
    from reportlab.lib.pagesizes import A4, LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, ListFlowable, ListItem, Image
    styles = getSampleStyleSheet()
    font = find_font(spec)
    text_blob = json.dumps(spec, ensure_ascii=False)
    if font:
        pdfmetrics.registerFont(TTFont('SenseFont', font))
        for s in styles.byName.values(): s.fontName = 'SenseFont'
        styles['BodyText'].fontSize = positive(spec.get('font_size', 10), 'font_size')
        styles['BodyText'].leading = styles['BodyText'].fontSize * 1.3
    elif any(ord(c) > 127 for c in text_blob):
        raise SpecError('non-Latin text needs a TTF font: pass "font_path" in the spec')
    esc = lambda s: html.escape(str(s)).replace('\n', '<br/>')
    story = []
    if spec.get('title'): story += [Paragraph(esc(spec['title']), styles['Title']), Spacer(1, 8)]
    if spec.get('subtitle'): story += [Paragraph(esc(spec['subtitle']), styles['Italic']), Spacer(1, 12)]
    for b in blocks_of(spec):
        t = b['type']
        if t == 'heading': story.append(Paragraph(esc(b['text']), styles['Heading%d' % max(1, min(int(b.get('level', 1)), 3))]))
        elif t == 'paragraph': story += [Paragraph(esc(b['text']), styles['BodyText']), Spacer(1, 6)]
        elif t in ('bullets', 'numbered'):
            story += [ListFlowable([ListItem(Paragraph(esc(i), styles['BodyText'])) for i in b.get('items', [])],
                                   bulletType='bullet' if t == 'bullets' else '1'), Spacer(1, 8)]
        elif t == 'table':
            data = [[Paragraph(esc(c), styles['BodyText']) for c in b['headers']]] + \
                   [[Paragraph(esc('' if c is None else c), styles['BodyText']) for c in r] for r in b['rows']]
            tab = Table(data, repeatRows=1)
            tab.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.5, colors.grey), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E7EEF7')),
                                     ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
            story += [tab, Spacer(1, 8)]
        elif t == 'page_break': story.append(PageBreak())
        elif t == 'image':
            if not Path(b['path']).is_file(): raise SpecError('image not found: ' + b['path'])
            image = Image(b['path'])
            max_w = min(positive(b.get('width_in', 6), 'width_in') * 72, 451)
            max_h = min(positive(b.get('height_in', 8), 'height_in') * 72, 620)
            ratio = min(max_w / image.imageWidth, max_h / image.imageHeight)
            image.drawWidth, image.drawHeight = image.imageWidth * ratio, image.imageHeight * ratio
            story.append(image)
    size = LETTER if str(spec.get('page_size', 'A4')).upper() == 'LETTER' else A4
    SimpleDocTemplate(out, pagesize=size, title=str(spec.get('title', '')), author=str(spec.get('author', ''))).build(story)


def inspect_pdf(src):
    need('pypdf', 'pypdf')
    from pypdf import PdfReader
    r = PdfReader(src)
    pages = [p.extract_text() or '' for p in r.pages]
    return {'format': 'pdf', 'pages': len(pages), 'text_by_page': pages,
            'note': 'Empty text on a page usually means a scanned image (OCR needed), not an empty page.' if any(not t.strip() for t in pages) else ''}


def merge_pdf(inputs, out):
    need('pypdf', 'pypdf')
    from pypdf import PdfWriter
    w = PdfWriter()
    for src in inputs: w.append(src)
    with open(out, 'wb') as f: w.write(f)
    return {'output': out, 'inputs': inputs}


def split_pdf(src, out_dir):
    need('pypdf', 'pypdf')
    from pypdf import PdfReader, PdfWriter
    r = PdfReader(src); Path(out_dir).mkdir(parents=True, exist_ok=True); outs = []
    for i, page in enumerate(r.pages, 1):
        w = PdfWriter(); w.add_page(page); p = str(Path(out_dir) / f'{Path(src).stem}-p{i:03d}.pdf')
        with open(p, 'wb') as f: w.write(f)
        outs.append(p)
    return {'outputs': outs}


# ---------------------------------------------------------------- XLSX
def create_xlsx(spec, out):
    object_spec(spec)
    need('openpyxl', 'openpyxl')
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    sheets = array(spec.get('sheets', []), 'sheets')
    if not sheets: raise SpecError('spec needs "sheets": [{"name": ..., "rows": [[...]]}]')
    wb = Workbook(); wb.remove(wb.active)
    used_names = set()
    for sh in sheets:
        object_spec(sh)
        name = str(sh.get('name', 'Sheet'))
        if not name or len(name) > 31 or re.search(r'[\\/*?:\[\]]', name) or name.startswith("'") or name.endswith("'"):
            raise SpecError('invalid sheet name: ' + name)
        if name.casefold() in used_names: raise SpecError('duplicate sheet name: ' + name)
        used_names.add(name.casefold())
        array(sh.get('rows', []), 'sheet rows')
        if not isinstance(sh.get('formulas', True), bool): raise SpecError('formulas must be true or false')
        if not isinstance(sh.get('number_formats', {}), dict) or not isinstance(sh.get('col_widths', {}), dict): raise SpecError('number_formats and col_widths must be objects')
        ws = wb.create_sheet(name)
        allow = sh.get('formulas', True)  # False: text starting with '=' is stored as plain text, never as a formula
        for r_idx, row in enumerate(sh.get('rows', []), 1):
            array(row, 'sheet row')
            for c_idx, v in enumerate(row, 1):
                if isinstance(v, (dict, list)): raise SpecError('cell values must be scalar')
                if isinstance(v, float) and not math.isfinite(v): raise SpecError('cell numbers must be finite')
                if isinstance(v, str) and len(v) > 32767: raise SpecError('cell text exceeds 32767 characters')
                cell = ws.cell(row=r_idx, column=c_idx, value=v)
                if isinstance(v, str) and v.startswith('=') and not allow:
                    cell.data_type = 's'
        if sh.get('header', True) and ws.max_row >= 1:
            for c in ws[1]:
                c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=sh.get('header_color', '1F4E79'))
                c.alignment = Alignment(vertical='center', wrap_text=True)
            ws.freeze_panes = 'A2'
            if ws.max_row > 1: ws.auto_filter.ref = ws.dimensions
        if sh.get('role') == 'inputs' or name.casefold() in ('inputs', 'input', 'ввод'):
            for row in ws.iter_rows(min_row=2 if sh.get('header', True) else 1):
                for cell in row:
                    if cell.data_type != 'f': cell.font = Font(color='0000FF')
        for col, fmt in sh.get('number_formats', {}).items():
            for cell in ws[col][1 if sh.get('header', True) else 0:]: cell.number_format = fmt
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            width = sh.get('col_widths', {}).get(letter)
            if width is None:
                width = min(60, max(10, max(len(str(c.value or '')) for c in ws[letter]) + 2))
            ws.column_dimensions[letter].width = positive(width, 'column width')
    wb.save(out)


def inspect_xlsx(src, max_rows=200):
    need('openpyxl', 'openpyxl')
    from openpyxl import load_workbook
    wb = load_workbook(src, data_only=False)
    out = {'format': 'xlsx', 'sheets': []}
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True, max_row=min(ws.max_row, max_rows))]
        formulas = sum(1 for r in ws.iter_rows() for c in r if c.data_type == 'f')
        out['sheets'].append({'name': ws.title, 'dimensions': ws.dimensions, 'rows': ws.max_row, 'columns': ws.max_column,
                              'formula_cells': formulas, 'preview': rows, 'truncated': ws.max_row > max_rows})
    out['note'] = 'Formulas are shown as text; values of formulas are not recalculated by this helper.'
    return out


# ---------------------------------------------------------------- PPTX
def create_pptx(spec, out):
    object_spec(spec)
    need('pptx', 'python-pptx')
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    theme = object_spec(spec.get('theme', {}))
    slides_spec = array(spec.get('slides', []), 'slides')
    for sd in slides_spec:
        object_spec(sd)
        if 'bullets' in sd: array(sd['bullets'], 'slide bullets')
        if sd.get('layout') == 'table': table_spec(sd)
        if 'font_size' in sd: positive(sd['font_size'], 'font_size')
    accent = RGBColor.from_string(theme.get('accent', '1F4E79')); font = theme.get('font', 'Calibri')
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height

    def style_title(shape, size=36):
        for p in shape.text_frame.paragraphs:
            for r in p.runs: r.font.size = Pt(size); r.font.bold = True; r.font.name = font; r.font.color.rgb = accent

    def text_box(slide, left, top, width, height, lines, size=20):
        tf = slide.shapes.add_textbox(left, top, width, height).text_frame; tf.word_wrap = True
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = ('• ' if len(lines) > 1 else '') + str(line); p.space_after = Pt(8)
            for r in p.runs: r.font.size = Pt(size); r.font.name = font
        return tf

    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    blank = prs.slide_layouts[6]

    def bar(slide, top, height=Inches(0.08), left=Inches(0.6), width=None):
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width or (W - Inches(1.2)), height)
        shp.fill.solid(); shp.fill.fore_color.rgb = accent; shp.line.fill.background()
        return shp

    def centered(slide, top, height, text, size, bold=True, color=None):
        tb = slide.shapes.add_textbox(Inches(0.8), top, W - Inches(1.6), height)
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.text = str(text)
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(size); r.font.bold = bold; r.font.name = font
                r.font.color.rgb = color or accent
        return tb

    if spec.get('title'):
        s = prs.slides.add_slide(blank)
        centered(s, Inches(2.3), Inches(1.5), spec['title'], 44)
        bar(s, Inches(3.95), left=Inches(4.7), width=W - Inches(9.4))
        if spec.get('subtitle'): centered(s, Inches(4.2), Inches(1), spec['subtitle'], 22, bold=False, color=RGBColor(0x40, 0x40, 0x40))
    expanded = []
    for sd in slides_spec:
        if sd.get('layout') == 'table' and len(sd.get('rows', [])) > 7:
            for offset in range(0, len(sd['rows']), 7):
                page = dict(sd, rows=sd['rows'][offset:offset+7])
                if offset: page['title'] = str(sd.get('title', '')) + ' (continued)'
                expanded.append(page)
        elif sd.get('layout', 'bullets') == 'bullets' and len(sd.get('bullets', [])) > 6:
            for offset in range(0, len(sd['bullets']), 6):
                page = dict(sd, bullets=sd['bullets'][offset:offset+6])
                if offset: page['title'] = str(sd.get('title', '')) + ' (continued)'
                expanded.append(page)
        else: expanded.append(sd)
    for sd in expanded:
        layout = sd.get('layout', 'bullets')
        s = prs.slides.add_slide(blank)
        if layout == 'section':
            band = bar(s, Inches(2.6), height=Inches(2.3), left=0, width=W)
            centered(s, Inches(2.6), Inches(2.3), sd.get('title', ''), 40, color=RGBColor(0xFF, 0xFF, 0xFF))
            if sd.get('notes'): s.notes_slide.notes_text_frame.text = str(sd['notes'])
            continue
        tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), W - Inches(1.2), Inches(1.1)); tb.text_frame.word_wrap = True
        tb.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
        tb.text_frame.text = str(sd.get('title', '')); style_title(tb, 30)
        bar(s, Inches(1.5))
        body_top, body_h = Inches(1.8), H - Inches(2.4)
        if layout == 'bullets':
            text_box(s, Inches(0.8), body_top, W - Inches(1.6), body_h, sd.get('bullets', []), sd.get('font_size', 22))
        elif layout == 'table':
            headers, rows = sd.get('headers', []), sd.get('rows', [])
            if not headers or any(len(r) != len(headers) for r in rows): raise SpecError('table slide: rows must match headers')
            t = s.shapes.add_table(len(rows) + 1, len(headers), Inches(0.6), body_top, W - Inches(1.2), Inches(0.55) * (len(rows) + 1)).table
            for i, r in enumerate([headers] + rows):
                for j, v in enumerate(r):
                    cell = t.cell(i, j); cell.text = '' if v is None else str(v)
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs: run.font.size = Pt(18); run.font.name = font; run.font.bold = (i == 0)
        elif layout == 'image':
            img = sd.get('image')
            if not img or not Path(img).is_file(): raise SpecError('image slide: file not found: ' + str(img))
            from PIL import Image as PILImage
            with PILImage.open(img) as im: iw, ih = im.size
            available_w = W - Inches(1.6)
            ratio = min(available_w / iw, body_h / ih)
            s.shapes.add_picture(img, int((W - iw * ratio) / 2), body_top, width=int(iw * ratio), height=int(ih * ratio))
        else:
            raise SpecError('unknown slide layout: ' + layout)
        if sd.get('notes'): s.notes_slide.notes_text_frame.text = str(sd['notes'])
    if not len(prs.slides): raise SpecError('no slides in spec')
    prs.save(out)


def inspect_pptx(src):
    need('pptx', 'python-pptx')
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    p = Presentation(src)
    slides = []
    for i, s in enumerate(p.slides, 1):
        texts, tables, charts, outside = [], [], [], []
        def visit(shapes):
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    visit(sh.shapes)
                if sh.has_text_frame and sh.text_frame.text.strip(): texts.append(sh.text_frame.text)
                if sh.has_table:
                    rows = [[c.text for c in row.cells] for row in sh.table.rows]
                    tables.append(rows)
                if sh.has_chart: charts.append(sh.name)
                if sh.left < 0 or sh.top < 0 or sh.left + sh.width > p.slide_width or sh.top + sh.height > p.slide_height:
                    outside.append(sh.name)
        visit(s.shapes)
        table_text = [cell for table in tables for row in table for cell in row]
        slides.append({'n': i, 'texts': texts, 'tables': tables, 'words': sum(len(t.split()) for t in texts + table_text),
                       'charts_requiring_visual_review': charts, 'outside_slide': outside,
                       'notes': s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''})
    return {'format': 'pptx', 'slides': slides, 'count': len(slides),
            'note': 'Includes table cells and grouped text. Chart content, image text and text-box overflow require rendering and visual review.'}


# ---------------------------------------------------------------- CLI
CREATORS = {'.docx': create_docx, '.pdf': create_pdf, '.xlsx': create_xlsx, '.pptx': create_pptx}
INSPECTORS = {'.docx': inspect_docx, '.pdf': inspect_pdf, '.xlsx': inspect_xlsx, '.pptx': inspect_pptx}


def doctor():
    return {'libraries': {k: probe_module(v) for k, v in LIBS.items()},
            'unicode_font_for_pdf': next((p for p in FONT_CANDIDATES if Path(p).is_file()), None)}


def main(argv=None):
    ap = argparse.ArgumentParser(description='Create, inspect, merge or split DOCX/PDF/XLSX/PPTX files locally')
    sub = ap.add_subparsers(dest='cmd', required=True)
    sub.add_parser('doctor')
    p = sub.add_parser('create'); p.add_argument('spec'); p.add_argument('output')
    p = sub.add_parser('inspect'); p.add_argument('file')
    p = sub.add_parser('merge'); p.add_argument('files', nargs='+')
    p = sub.add_parser('split'); p.add_argument('file'); p.add_argument('out_dir')
    a = ap.parse_args(argv)
    try:
        if a.cmd == 'doctor': res = doctor()
        elif a.cmd == 'create':
            ext = Path(a.output).suffix.lower()
            if ext not in CREATORS: raise SpecError('output must end with .docx, .pdf, .xlsx or .pptx')
            spec = json.loads(Path(a.spec).read_text(encoding='utf-8'))
            Path(a.output).parent.mkdir(parents=True, exist_ok=True)
            fd, temp_path = tempfile.mkstemp(suffix=ext, dir=Path(a.output).parent)
            os.close(fd)
            try:
                CREATORS[ext](spec, temp_path)
                os.replace(temp_path, a.output)
            finally:
                if os.path.exists(temp_path): os.unlink(temp_path)
            res = {'output': a.output, 'bytes': os.path.getsize(a.output), 'next': f'verify with: python scripts/docs.py inspect {a.output}'}
        elif a.cmd == 'inspect':
            ext = Path(a.file).suffix.lower()
            if ext not in INSPECTORS: raise SpecError('unsupported file type: ' + ext)
            res = INSPECTORS[ext](a.file)
        elif a.cmd == 'merge':
            if len(a.files) < 3: raise SpecError('merge needs at least two input PDFs and an output path')
            res = merge_pdf(a.files[:-1], a.files[-1])
        else:
            res = split_pdf(a.file, a.out_dir)
        print(json.dumps(res, ensure_ascii=False, indent=2, default=str))
        return 0
    except Exception as e:
        print('error: ' + str(e), file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
